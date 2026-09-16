"""
중간보고서_v2.md → 중간보고서_v2.docx 변환 스크립트
- 제목/소제목 스타일링
- 테이블 변환
- 코드블록 변환
- 한글 폰트(맑은 고딕) 적용
"""
import re
from pathlib import Path
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

ROOT = Path("/Users/joaram/Desktop/SCHU/졸업작품/ver1/docs")
SRC = ROOT / "중간보고서_v2.md"
DST = ROOT / "중간보고서_v2.docx"

KOR_FONT = "맑은 고딕"
ENG_FONT = "Times New Roman"
MONO_FONT = "Consolas"


def set_run_font(run, font_name=KOR_FONT, size=11, bold=False, color=None):
    run.font.name = font_name
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = color
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    rFonts.set(qn("w:eastAsia"), KOR_FONT)
    rFonts.set(qn("w:ascii"), font_name)
    rFonts.set(qn("w:hAnsi"), font_name)


def set_cell_shading(cell, color_hex):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), color_hex)
    tc_pr.append(shd)


# 인라인 마크다운 처리: **bold**, *italic*, `code`
INLINE_RE = re.compile(r"(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`)")


def add_inline(paragraph, text, base_size=11, base_bold=False):
    parts = INLINE_RE.split(text)
    for part in parts:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            set_run_font(run, size=base_size, bold=True)
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            run = paragraph.add_run(part[1:-1])
            set_run_font(run, size=base_size, bold=base_bold)
            run.italic = True
        elif part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            set_run_font(run, font_name=MONO_FONT, size=base_size - 1, bold=False)
        else:
            run = paragraph.add_run(part)
            set_run_font(run, size=base_size, bold=base_bold)


def add_heading(doc, text, level):
    sizes = {1: 20, 2: 16, 3: 14, 4: 12, 5: 11, 6: 11}
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12 if level > 1 else 18)
    p.paragraph_format.space_after = Pt(6)
    if level == 1:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    set_run_font(run, size=sizes.get(level, 11), bold=True,
                 color=RGBColor(0x1F, 0x3A, 0x5F) if level <= 2 else RGBColor(0x33, 0x33, 0x33))


def add_paragraph(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.5
    add_inline(p, text)


def add_list_item(doc, text, level=0):
    p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    p.paragraph_format.left_indent = Cm(0.7 + level * 0.5)
    p.paragraph_format.space_after = Pt(2)
    add_inline(p, text)


def add_numbered_item(doc, text):
    p = doc.add_paragraph(style="List Number")
    p.paragraph_format.space_after = Pt(2)
    add_inline(p, text)


def add_table(doc, rows):
    """rows: list[list[str]] including header at index 0"""
    if not rows:
        return
    ncols = max(len(r) for r in rows)
    rows = [r + [""] * (ncols - len(r)) for r in rows]
    table = doc.add_table(rows=len(rows), cols=ncols)
    table.style = "Light Grid Accent 1"
    table.autofit = True
    for i, row in enumerate(rows):
        for j, cell_text in enumerate(row):
            cell = table.cell(i, j)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            cell.text = ""
            p = cell.paragraphs[0]
            add_inline(p, cell_text.strip(), base_size=10, base_bold=(i == 0))
            if i == 0:
                set_cell_shading(cell, "1F3A5F")
                for run in p.runs:
                    run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    doc.add_paragraph()


def add_code_block(doc, code_lines):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.5)
    p.paragraph_format.space_after = Pt(6)
    pPr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), "F5F5F5")
    pPr.append(shd)
    text = "\n".join(code_lines)
    run = p.add_run(text)
    set_run_font(run, font_name=MONO_FONT, size=9)


def parse_markdown(md_text):
    lines = md_text.split("\n")
    doc = Document()

    # 기본 마진 / 페이지 설정
    for section in doc.sections:
        section.top_margin = Cm(2.2)
        section.bottom_margin = Cm(2.2)
        section.left_margin = Cm(2.5)
        section.right_margin = Cm(2.5)

    # 기본 폰트
    style = doc.styles["Normal"]
    style.font.name = KOR_FONT
    style.font.size = Pt(11)
    rPr = style.element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    rFonts.set(qn("w:eastAsia"), KOR_FONT)
    rFonts.set(qn("w:ascii"), KOR_FONT)
    rFonts.set(qn("w:hAnsi"), KOR_FONT)

    i = 0
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        # 코드블록
        if stripped.startswith("```"):
            i += 1
            buf = []
            while i < len(lines) and not lines[i].strip().startswith("```"):
                buf.append(lines[i])
                i += 1
            add_code_block(doc, buf)
            i += 1
            continue

        # 테이블
        if "|" in line and i + 1 < len(lines) and re.match(r"^\s*\|?[\s:|-]+\|?\s*$", lines[i + 1]):
            tbl = []
            while i < len(lines) and "|" in lines[i]:
                row_line = lines[i].strip()
                if re.match(r"^\s*\|?[\s:|-]+\|?\s*$", row_line):
                    i += 1
                    continue
                cells = [c.strip() for c in row_line.strip("|").split("|")]
                tbl.append(cells)
                i += 1
            add_table(doc, tbl)
            continue

        # 헤딩
        m = re.match(r"^(#{1,6})\s+(.*)", stripped)
        if m:
            level = len(m.group(1))
            add_heading(doc, m.group(2), level)
            i += 1
            continue

        # 수평선
        if stripped == "---":
            doc.add_paragraph().add_run("―" * 30)
            i += 1
            continue

        # 빈 줄
        if not stripped:
            i += 1
            continue

        # 리스트
        m_ul = re.match(r"^(\s*)[-*+]\s+(.*)", line)
        m_ol = re.match(r"^(\s*)\d+\.\s+(.*)", line)
        if m_ul:
            indent_level = len(m_ul.group(1)) // 2
            add_list_item(doc, m_ul.group(2), level=min(indent_level, 1))
            i += 1
            continue
        if m_ol:
            add_numbered_item(doc, m_ol.group(2))
            i += 1
            continue

        # 일반 문단
        add_paragraph(doc, stripped)
        i += 1

    return doc


def main():
    md_text = SRC.read_text(encoding="utf-8")
    doc = parse_markdown(md_text)
    doc.save(DST)
    print(f"✅ 저장 완료: {DST}")
    print(f"   파일 크기: {DST.stat().st_size / 1024:.1f} KB")


if __name__ == "__main__":
    main()
